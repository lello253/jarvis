import os
import json
from pathlib import Path
from PyQt6.QtWidgets import QMenu, QMessageBox
from PyQt6.QtGui import QCursor

def process_dropped_file(file_path: str, parent_window):
    """Gestisce il rilascio del file sull'HUD azionando comandi reali e salvataggio nel Vault."""
    path = Path(file_path)
    if not path.exists():
        return

    ext = path.suffix.lower()
    
    parent_window._active_drop_menu = QMenu(parent_window)
    menu = parent_window._active_drop_menu

    menu.setStyleSheet("""
        QMenu {
            background-color: #0b0e14;
            border: 1px solid #00f0ff;
            border-radius: 8px;
            padding: 6px;
        }
        QMenu::item {
            padding: 8px 18px;
            color: #c9d1d9;
            font-family: 'Segoe UI';
            font-size: 12px;
        }
        QMenu::item:selected {
            background-color: #161b22;
            color: #00f0ff;
        }
    """)

    header = menu.addAction(f"📄 FILE: {path.name.upper()}")
    header.setEnabled(False)
    menu.addSeparator()

    if ext in ('.py', '.js', '.cpp', '.c', '.cs', '.html', '.css', '.json'):
        act_debug = menu.addAction("🐛 Analizza & Trova Errori")
        act_explain = menu.addAction("📝 Spiega Codice")
        act_save = menu.addAction("💾 Salva nel Memory Vault")

        act_debug.triggered.connect(lambda: _read_and_send(parent_window, path, "Analizza questo codice e trova eventuali errori o bug:"))
        act_explain.triggered.connect(lambda: _read_and_send(parent_window, path, "Spiegami in modo sintetico la logica di questo codice:"))
        act_save.triggered.connect(lambda: _copy_to_vault(path, parent_window))

    elif ext in ('.txt', '.md', '.log', '.pdf', '.docx'):
        act_sum = menu.addAction("⚡ Riassumi in 3 punti")
        act_vault = menu.addAction("🧠 Estrai Testo e Salva nel Vault")

        act_sum.triggered.connect(lambda: _read_and_send(parent_window, path, "Fai un riassunto sintetico in 3 punti del contenuto di questo file:"))
        act_vault.triggered.connect(lambda: _copy_to_vault(path, parent_window))

    elif ext in ('.png', '.jpg', '.jpeg', '.webp'):
        act_ocr = menu.addAction("👁️ Descrivi / Analizza Immagine")
        act_ocr.triggered.connect(lambda: _send_prompt(parent_window, f"Analizza l'immagine trascinata al percorso: {path.absolute()}"))

    else:
        act_gen = menu.addAction("📥 Salva nel Memory Vault")
        act_gen.triggered.connect(lambda: _copy_to_vault(path, parent_window))

    menu.popup(QCursor.pos())

def _send_prompt(parent, prompt_text: str):
    """Invia il prompt a JARVIS provando le callback della finestra principale."""
    try:
        if hasattr(parent, 'on_text_command') and callable(parent.on_text_command):
            parent.on_text_command(prompt_text)
            return

        for send_method_name in ('_send_text_command', '_submit_prompt', 'send_command', '_on_send'):
            if hasattr(parent, send_method_name):
                method = getattr(parent, send_method_name)
                if callable(method):
                    method(prompt_text)
                    return
    except Exception as e:
        print(f"[SmartDrop Input Error] {e}")

def _read_and_send(parent, path: Path, instruction: str):
    """Legge il testo/documento e lo inoltra al motore principale di JARVIS."""
    try:
        content = ""
        ext = path.suffix.lower()

        if ext in ('.txt', '.md', '.log', '.json', '.py', '.js', '.cpp', '.c', '.html', '.css'):
            content = path.read_text(encoding='utf-8', errors='ignore')
        elif ext == '.docx':
            try:
                import docx
                doc = docx.Document(path)
                content = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            except Exception:
                content = f"File Word: {path.name}"
        elif ext == '.pdf':
            try:
                import pypdf
                reader = pypdf.PdfReader(path)
                content = "\n\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            except Exception:
                content = f"File PDF: {path.name}"
        else:
            content = path.read_text(encoding='utf-8', errors='ignore')

        prompt = f"{instruction}\n\n{content[:4000]}"
        _send_prompt(parent, prompt)
    except Exception as e:
        QMessageBox.warning(parent, "Errore Lettura", f"Impossibile leggere il file: {e}")

def _get_vault_dir(parent) -> Path:
    config_path = Path("./config/settings.json")
    if config_path.exists():
        try:
            data = json.loads(config_path.read_text(encoding="utf-8"))
            if "memory_folder" in data and data["memory_folder"]:
                p = Path(data["memory_folder"])
                p.mkdir(parents=True, exist_ok=True)
                return p
        except Exception:
            pass

    for attr in ('vault_path', 'memory_dir', '_memory_dir'):
        if hasattr(parent, attr):
            val = getattr(parent, attr)
            if val:
                p = Path(val)
                p.mkdir(parents=True, exist_ok=True)
                return p

    fallback = Path("./app_memory")
    fallback.mkdir(parents=True, exist_ok=True)
    return fallback

def _copy_to_vault(file_path: Path, parent):
    try:
        vault_dir = _get_vault_dir(parent)
        extracted_text = ""
        ext = file_path.suffix.lower()

        if ext in ('.txt', '.md', '.log', '.json', '.py', '.js', '.cpp', '.c', '.html', '.css'):
            extracted_text = file_path.read_text(encoding='utf-8', errors='ignore')
        elif ext == '.docx':
            try:
                import docx
                doc = docx.Document(file_path)
                extracted_text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            except ImportError:
                extracted_text = f"*(Modulo python-docx non installato. Percorso: `{file_path.absolute()}`)*"
        elif ext == '.pdf':
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                pages_text = [page.extract_text() for page in reader.pages if page.extract_text()]
                extracted_text = "\n\n--- Interruzione Pagina ---\n\n".join(pages_text)
            except ImportError:
                extracted_text = f"*(Modulo pypdf non installato. Percorso: `{file_path.absolute()}`)*"
        else:
            extracted_text = f"File sorgente: `{file_path.name}`\nPercorso: `{file_path.absolute()}`"

        md_content = f"# {file_path.stem}\n\n> **Nota importata da JARVIS Smart Drop**\n> **Sorgente:** `{file_path.name}`\n\n---\n\n{extracted_text}"

        dest_file = vault_dir / f"{file_path.stem}.md"
        dest_file.write_text(md_content, encoding='utf-8')

        if hasattr(parent, 'vault_window') and parent.vault_window:
            if hasattr(parent.vault_window, 'refresh_vault'):
                parent.vault_window.refresh_vault()
            elif hasattr(parent.vault_window, 'model'):
                parent.vault_window.model.setRootPath(str(vault_dir))

        QMessageBox.information(parent, "Memory Vault", f"Nota creata con successo!\n\n📄 {dest_file.name}")
    except Exception as e:
        QMessageBox.critical(parent, "Errore Vault", f"Impossibile creare la nota: {e}")